import base64
import json
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch

from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import book_import
import epub_export
from book_import import import_book, safe_target, epub_metadata, defaults

PIXEL = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jv1cAAAAASUVORK5CYII="
)


class Adapters(unittest.TestCase):
    def test_html_encoding_roundtrips_utf8_and_declared_legacy_text(self):
        for encoding, meta in [("utf-8", ""), ("gb18030", '<meta charset="gb18030"/>')]:
            source = book_import.BOOK_ROOT / (encoding + '.html')
            source.parent.mkdir(parents=True, exist_ok=True)
            source.write_bytes(('<html lang="en"><head>' + meta + '<title>原创 σ</title></head><body><h1>第一章</h1><p>映射 τ 与集合。</p></body></html>').encode(encoding))
            metadata = book_import.html_metadata(source)
            self.assertEqual(metadata['title'], '原创 σ')
            self.assertEqual(metadata['toc'][0]['title'], '第一章')
            self.assertEqual(metadata['language'], 'en')
            restored = book_import.load_html(source)
            self.assertIn('映射 τ 与集合。', ''.join(restored.getroot().itertext()))
            if meta:
                self.assertEqual(restored.xpath('string(//meta/@charset)'), 'utf-8')
    def test_embedded_data_images_are_real_resources_in_html_and_markdown_imports(self):
        samples = {
            "png": ("image/png;base64," + base64.b64encode(PIXEL).decode(), PIXEL),
            "svg": (
                "image/svg+xml,%3Csvg%20xmlns=%22http://www.w3.org/2000/svg%22%20viewBox=%220%200%2010%2010%22%3E%3Ccircle%20cx=%225%22%20cy=%225%22%20r=%224%22/%3E%3C/svg%3E",
                b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><circle cx="5" cy="5" r="4"/></svg>',
            ),
        }
        for variant in ("html", "md", "raw-block", "raw-inline"):
            extension = "html" if variant == "html" else "md"
            for kind, (uri, expected) in samples.items():
                with self.subTest(format=variant, image=kind):
                    source = self.root / (
                        "image-" + kind + "-" + variant + "." + extension
                    )
                    source.write_text(
                        '<h1>Original image sample</h1><p>Original text.</p><img src="data:'
                        + uri
                        + '" alt="Original illustration"/>'
                        if extension == "html"
                        else "# Original image sample\n\nOriginal text.\n\n![Original illustration](data:"
                        + uri
                        + ")"
                    )
                    if variant in ("raw-block", "raw-inline"):
                        tag = (
                            '<img src="data:' + uri + '" alt="Original illustration"/>'
                        )
                        source.write_text(
                            "# Original raw image\n\n"
                            + (
                                "<div>" + tag + "</div>"
                                if variant == "raw-block"
                                else "Original inline image "
                                + tag
                                + " and remaining text."
                            )
                        )
                    book = import_book(source)
                    content = book_import.BOOK_ROOT / book["id"] / "content"
                    assets = [
                        path
                        for path in content.rglob("*" + "." + kind)
                        if not path.name.startswith("._")
                    ]
                    self.assertEqual(len(assets), 1)
                    self.assertEqual(assets[0].read_bytes(), expected)
                    html = "".join(
                        book_import.path_for(chapter["url"]).read_text()
                        for chapter in book["chapters"]
                    )
                    self.assertNotIn("../data:image", html)
                    self.assertEqual(
                        book_import.path_for(book["originalUrl"]).read_bytes(),
                        source.read_bytes(),
                    )

    def test_external_images_are_not_fetched_and_failed_conversion_preserves_the_target(
        self,
    ):
        requests = []

        class Provider(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_GET(self):
                requests.append(self.path)
                self.send_response(200)
                self.end_headers()
                self.wfile.write(PIXEL)

        provider = ThreadingHTTPServer(("127.0.0.1", 0), Provider)
        worker = threading.Thread(target=provider.serve_forever, daemon=True)
        worker.start()
        target = self.root / "previous.epub"
        target.write_bytes(b"previous completed output")
        source = self.root / "external.md"
        try:
            for address in (
                f"http://127.0.0.1:{provider.server_port}/private.png",
                "private.png",
            ):
                (self.root / "private.png").write_bytes(PIXEL)
                source.write_text(
                    "# Original sample\n\n![Unavailable asset](" + address + ")"
                )
                with self.assertRaisesRegex(ValueError, "图片未包含"):
                    book_import.convert_document(source, target)
                self.assertEqual(target.read_bytes(), b"previous completed output")
            source.write_text(
                "# Original sample\n\n![Invalid](data:text/plain;base64,cHJpdmF0ZQ==)"
            )
            with self.assertRaisesRegex(ValueError, "内嵌图片"):
                book_import.convert_document(source, target)
            self.assertEqual(requests, [])
            source.write_text(
                '# Original raw block\n\n<div><img src="data:image/png;base64,'
                + base64.b64encode(PIXEL).decode()
                + '"/><iframe src="http://127.0.0.1:'
                + str(provider.server_port)
                + '/private.png"></iframe></div>'
            )
            book_import.convert_document(source, self.root / "isolated.epub")
            self.assertEqual(requests, [])
            self.assertEqual(target.read_bytes(), b"previous completed output")
        finally:
            provider.shutdown()
            provider.server_close()
            worker.join()

    def setUp(self):
        (book_import.ROOT / "tmp").mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(
            prefix="adapter-tests-",
            dir=book_import.ROOT / "tmp",
            ignore_cleanup_errors=True,
        )
        self.root = Path(self.temp.name)
        self.override = patch.object(book_import, "BOOK_ROOT", self.root / "books")
        self.override.start()

    def tearDown(self):
        self.override.stop()
        self.temp.cleanup()

    def epub(self, navigation=True):
        path = self.root / ("nav.epub" if navigation else "ncx.epub")
        files = {
            "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>',
            "OPS/book.opf": '<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>通用格式验证</dc:title><dc:creator>测试作者</dc:creator><dc:language>zh</dc:language></metadata><manifest><item id="one" href="one.xhtml" media-type="application/xhtml+xml"/><item id="two" href="two.xhtml" media-type="application/xhtml+xml"/><item id="image" href="pic.png" media-type="image/png" properties="cover-image"/><item id="nav" href="nav.xhtml" properties="nav" media-type="application/xhtml+xml"/><item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest><spine toc="ncx"><itemref idref="one"/><itemref idref="two"/></spine></package>',
            "OPS/one.xhtml": '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>第一章</title></head><body><h1 id="start">第一章</h1><p>正文<em>强调</em>。<img src="pic.png"/></p><a href="two.xhtml#note">注释</a></body></html>',
            "OPS/two.xhtml": '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>第二章</title></head><body><h1>第二章</h1><p id="note">脚注正文</p></body></html>',
            "OPS/nav.xhtml": '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>目录</title></head><body><nav epub:type="toc"><ol><li><a href="one.xhtml#start">第一章</a><ol><li><a href="two.xhtml#note">脚注</a></li></ol></li><li><a href="two.xhtml">第二章</a></li></ol></nav></body></html>',
            "OPS/toc.ncx": '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/"><navMap><navPoint><navLabel><text>第一章</text></navLabel><content src="one.xhtml#start"/></navPoint><navPoint><navLabel><text>脚注</text></navLabel><content src="two.xhtml#note"/></navPoint></navMap></ncx>',
            "OPS/pic.png": PIXEL,
        }
        if not navigation:
            files["OPS/book.opf"] = files["OPS/book.opf"].replace(
                'properties="nav"', ""
            )
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("mimetype", "application/epub+zip")
            for name, value in files.items():
                archive.writestr(name, value)
        return path, files

    def test_epub3_preserves_resources_and_fragment_navigation(self):
        source, files = self.epub()
        book = import_book(source)
        self.assertFalse(book["parseEnabled"])
        self.assertEqual(book["tocSource"], "epub-nav")
        self.assertEqual(book["toc"][0]["children"][0]["index"], 2)
        self.assertEqual(book["toc"][0]["children"][0]["fragment"], "note")
        content = book_import.BOOK_ROOT / book["id"] / "content"
        self.assertEqual(
            (content / "OPS/one.xhtml").read_text(), files["OPS/one.xhtml"]
        )
        self.assertEqual((content / "OPS/pic.png").read_bytes(), PIXEL)
        self.assertEqual(len(book["chapters"]), 2)

    def test_epub2_reads_native_ncx(self):
        source, _ = self.epub(False)
        book = import_book(source)
        self.assertEqual(book["tocSource"], "epub-ncx")
        self.assertEqual(book["toc"][1]["fragment"], "note")

    def kindle(self, extension):
        source = self.root / ("sample" + extension)
        source.write_bytes(b"Original anonymous Kindle input" + extension.encode())
        unpacked = self.root / ("unpacked" + extension)
        unpacked.mkdir()
        body = '<head><title>Anonymous format sample</title></head><body><a id="part-a"></a><p>Alpha <em>emphasis</em> <img src="pic.png" alt="Original image"/></p><div height="1em"></div><a id="part-b"></a><p>Beta</p></body>'
        files = {
            "content.opf": '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="uid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:opf="http://www.idpf.org/2007/opf"><dc:identifier id="uid" opf:scheme="URI">urn:page-between:kindle-sample</dc:identifier><dc:title>Anonymous format sample</dc:title><dc:creator opf:role="aut">Original example</dc:creator><dc:language>en</dc:language><meta name="cover" content="cover"/></metadata><manifest><item id="body" href="book.html" media-type="text/html"/><item id="cover" href="pic.png" media-type="image/png"/><item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest><spine toc="ncx"><itemref idref="body"/></spine><tours/><guide><reference type="text" href="book.html#missing-start"/></guide></package>',
            "toc.ncx": '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head><meta name="dtb:uid" content="urn:page-between:kindle-sample"/></head><docTitle><text>Anonymous format sample</text></docTitle><navMap><navPoint id="a" playOrder="1"><navLabel><text>Part A</text></navLabel><content src="book.html#part-a"/><navPoint id="b" playOrder="2"><navLabel><text>Part B</text></navLabel><content src="book.html#part-b"/></navPoint></navPoint></navMap></ncx>',
            "book.html": '<html lang="en">' + body + "</html>",
            "pic.png": PIXEL,
        }
        if extension == ".azw3":
            files["book.html"] = (
                '<html xmlns="http://www.w3.org/1999/xhtml" lang="en">'
                + body.replace(
                    '<div height="1em"></div>', "<section><p>Modern XHTML</p></section>"
                )
                + "</html>"
            )
            files["content.opf"] = files["content.opf"].replace(
                'media-type="text/html"', 'media-type="application/xhtml+xml"'
            )
            files["META-INF/container.xml"] = (
                '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>'
            )
            extracted = unpacked / "book.epub"
            with zipfile.ZipFile(extracted, "w") as archive:
                archive.writestr("mimetype", "application/epub+zip")
                for name, value in files.items():
                    archive.writestr(name, value)
        else:
            for name, value in files.items():
                (unpacked / name).write_bytes(
                    value if isinstance(value, bytes) else value.encode()
                )
            extracted = unpacked / "book.html"
        return source, unpacked, extracted

    def test_kindle_preserves_native_directory_images_and_exports_valid_epub3(self):
        for extension in (".mobi", ".azw3"):
            with self.subTest(extension=extension):
                source, unpacked, extracted = self.kindle(extension)
                with patch(
                    "mobi.extract", return_value=(str(unpacked), str(extracted))
                ):
                    book = import_book(source)
                self.assertFalse(unpacked.exists())
                self.assertEqual(book["language"], "en")
                self.assertEqual(book["author"], "Original example")
                self.assertEqual(book["tocSource"], "epub-nav")
                self.assertEqual(book["chapters"][0]["title"], "Part A")
                self.assertEqual(book["toc"][0]["fragment"], "part-a")
                self.assertEqual(book["toc"][0]["children"][0]["fragment"], "part-b")
                self.assertEqual(book["toc"][0]["children"][0]["index"], 1)
                self.assertEqual(
                    book_import.path_for(book["cover"]).read_bytes(), PIXEL
                )
                self.assertEqual(
                    book_import.path_for(book["originalUrl"]).read_bytes(),
                    source.read_bytes(),
                )
                document = etree.parse(
                    str(book_import.path_for(book["chapters"][0]["url"]))
                )
                self.assertEqual(
                    len(document.xpath('//*[@id="part-a" or @id="part-b"]')), 2
                )
                self.assertTrue(
                    document.xpath('//*[local-name()="em"][text()="emphasis"]')
                )
                if extension == ".mobi":
                    self.assertTrue(
                        document.xpath('//*[local-name()="div"][@style="height:1em"]')
                    )
                with (
                    patch.object(epub_export, "BOOK_ROOT", book_import.BOOK_ROOT),
                    patch.object(epub_export, "CACHE", self.root / "cache"),
                ):
                    result = epub_export.build(book)
                    output = self.root / "cache/exports" / result["id"]
                    if checker := os.environ.get("EPUBCHECK_JAR"):
                        checked = subprocess.run(
                            ["java", "-jar", checker, str(output), "--failonwarnings"],
                            capture_output=True,
                            text=True,
                            timeout=60,
                        )
                        self.assertEqual(
                            checked.returncode, 0, checked.stdout + checked.stderr
                        )

    def test_kindle_directory_cannot_reference_files_outside_its_package(self):
        source, unpacked, extracted = self.kindle(".mobi")
        navigation = unpacked / "toc.ncx"
        navigation.write_text(
            navigation.read_text().replace("book.html#part-a", "../private.html#part-a")
        )
        with patch("mobi.extract", return_value=(str(unpacked), str(extracted))):
            with self.assertRaisesRegex(ValueError, "Kindle 文件无法导入"):
                import_book(source)
        self.assertFalse(unpacked.exists())

    def test_kindle_ambiguous_packages_do_not_silently_drop_the_native_directory(self):
        source, unpacked, extracted = self.kindle(".mobi")
        (unpacked / "other.opf").write_bytes((unpacked / "content.opf").read_bytes())
        with patch("mobi.extract", return_value=(str(unpacked), str(extracted))):
            with self.assertRaisesRegex(ValueError, "Kindle 文件无法导入"):
                import_book(source)
        self.assertFalse(unpacked.exists())

    def test_zip_cannot_escape_import_directory(self):
        source = self.root / "bad.epub"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("../escaped.txt", "bad")
        with self.assertRaises(ValueError):
            import_book(source)
        self.assertFalse((self.root / "escaped.txt").exists())
        with self.assertRaises(ValueError):
            safe_target(self.root, "https://example.com/book", self.root)

    def test_txt_switch_and_encoding(self):
        source = self.root / "章节.txt"
        source.write_bytes(
            "第一卷 春\n第一章 来信\n正文一\n第二章 出发\n正文二".encode("gb18030")
        )
        book = import_book(source)
        self.assertTrue(book["parseEnabled"])
        self.assertEqual(len(book["chapters"]), 2)
        self.assertEqual(book["toc"][0]["title"], "第一卷 春")
        plain = import_book(source, enabled=False, ident="plain")
        self.assertEqual(len(plain["chapters"]), 1)
        self.assertEqual(plain["tocSource"], "text")
        self.assertFalse(defaults(".epub"))
        self.assertTrue(defaults(".pdf"))

    def test_txt_offsets_use_the_same_normalized_source_with_parse_on_and_off(self):
        source = self.root / "windows.txt"
        text = "  \r\n第一章 来信\r\n正文一\r\n第二章 出发\r\n正文二"
        source.write_bytes(text.encode("utf-8-sig"))
        normalized = text.replace("\r\n", "\n")
        parsed = import_book(source)
        plain = import_book(source, enabled=False, ident="whole")
        self.assertEqual(
            parsed["chapters"][1]["sourceOffset"], normalized.index("正文二")
        )
        self.assertEqual(
            plain["chapters"][0]["sourceOffset"], normalized.index("第一章")
        )
        content = book_import.path_for(plain["contentUrl"])
        self.assertNotIn("\r", json.loads(content.read_text())[0]["body"])

    def test_docx_keeps_embedded_image(self):
        from docx import Document

        image = self.root / "pixel.png"
        image.write_bytes(PIXEL)
        source = self.root / "图文教材.docx"
        doc = Document()
        doc.add_heading("第一章 原文", level=1)
        doc.add_paragraph("保留正文与插图。")
        doc.add_picture(str(image))
        doc.save(source)
        book = import_book(source)
        self.assertEqual(book["type"], "epub")
        self.assertFalse(book["parseEnabled"])
        content = book_import.BOOK_ROOT / book["id"] / "content"
        self.assertTrue(list(content.rglob("*.png")))

    def test_cbz_retains_numeric_page_order(self):
        source = self.root / "漫画.cbz"
        with zipfile.ZipFile(source, "w") as archive:
            for name in ["10.png", "2.png", "1.png"]:
                archive.writestr(name, PIXEL)
        book = import_book(source)
        self.assertEqual(book["layout"], "comic")
        self.assertEqual(len(book["chapters"]), 3)
        first = book_import.path_for(book["chapters"][0]["url"])
        self.assertIn("1.png", first.read_text())

    def test_other_common_formats(self):
        samples = {
            ".md": "# 第一章\n\n正文 **加粗**。",
            ".html": '<!doctype html><html><head><meta charset="utf-8"><title>HTML 试读</title></head><body><h1>第一章</h1><p>原文。</p></body></html>',
            ".rtf": r"{\rtf1\ansi\deff0 Reading sample\par Second paragraph}",
            ".fb2": '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0"><description><title-info><book-title>FB2 试读</book-title><lang>zh</lang></title-info></description><body><section><title><p>第一章</p></title><p>正文。</p></section></body></FictionBook>',
        }
        for extension, text in samples.items():
            source = self.root / ("sample" + extension)
            source.write_text(text)
            with self.subTest(extension=extension):
                book = import_book(source)
                self.assertEqual(book["type"], "epub")
                self.assertGreater(len(book["chapters"]), 0)


if __name__ == "__main__":
    unittest.main()
